"""KB loaders + /knowledge-bases/{id}/sources/upload integration."""

from __future__ import annotations

import io

import pytest

from app.kb.loaders import detect_kind, extract_text


def test_detect_kind_basic():
    assert detect_kind("notes.txt") == "txt"
    assert detect_kind("readme.MD") == "md"
    assert detect_kind("Manual.pdf") == "pdf"
    assert detect_kind("doc.docx") == "docx"


def test_detect_kind_rejects_unknown():
    from app.kb.loaders import UnsupportedSourceError

    with pytest.raises(UnsupportedSourceError):
        detect_kind("video.mp4")


def test_extract_text_txt():
    out = extract_text(b"hello world\nsecond line", kind="txt")
    assert "hello world" in out


def test_extract_text_md():
    out = extract_text(b"# heading\n\nbody body body", kind="md")
    assert "heading" in out and "body" in out


def test_extract_text_docx_roundtrip(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("Acme refund policy")
    doc.add_paragraph("Refunds processed in 3-5 business days.")
    buf = io.BytesIO()
    doc.save(buf)
    out = extract_text(buf.getvalue(), kind="docx")
    assert "refund policy" in out.lower()
    assert "3-5 business days" in out


# ---------- upload endpoint -------------------------------------------------


@pytest.mark.asyncio
async def test_upload_txt_ingests_and_chunks(client, auth_headers, monkeypatch):
    # Patch the LLM embed call to avoid hitting Gemini/OpenAI in tests.
    async def fake_embed(*, model_id: str, texts: list[str]) -> list[list[float]]:
        return [[0.0] * 1536 for _ in texts]

    from app.kb import store as kb_store

    monkeypatch.setattr(kb_store, "embed", fake_embed)

    r = await client.post(
        "/v1/knowledge-bases",
        json={"name": "Policies"},
        headers=auth_headers,
    )
    assert r.status_code == 201, r.text
    kb_id = r.json()["id"]

    content = b"Section 1.\n\nLong policy text. " * 60  # multi-chunk
    files = {"file": ("policy.txt", content, "text/plain")}
    r = await client.post(
        f"/v1/knowledge-bases/{kb_id}/sources/upload",
        files=files,
        headers=auth_headers,
    )
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["kind"] == "txt"
    assert body["status"] == "ready"


@pytest.mark.asyncio
async def test_upload_unsupported_kind_415(client, auth_headers):
    r = await client.post("/v1/knowledge-bases", json={"name": "X"}, headers=auth_headers)
    kb_id = r.json()["id"]
    files = {"file": ("song.mp3", b"\x00\x01\x02", "audio/mpeg")}
    r = await client.post(
        f"/v1/knowledge-bases/{kb_id}/sources/upload",
        files=files,
        headers=auth_headers,
    )
    assert r.status_code == 415, r.text


@pytest.mark.asyncio
async def test_upload_empty_file_400(client, auth_headers):
    r = await client.post("/v1/knowledge-bases", json={"name": "Y"}, headers=auth_headers)
    kb_id = r.json()["id"]
    files = {"file": ("empty.txt", b"", "text/plain")}
    r = await client.post(
        f"/v1/knowledge-bases/{kb_id}/sources/upload",
        files=files,
        headers=auth_headers,
    )
    assert r.status_code == 400, r.text
