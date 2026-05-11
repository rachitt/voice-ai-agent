"""Source file → plain text loaders for KB ingestion.

Only stdlib + pypdf + python-docx. Streaming readers operate on in-memory
`bytes` since uploads are bounded by FastAPI request limits anyway.
"""
from __future__ import annotations

import io
from typing import Literal

Kind = Literal["txt", "md", "pdf", "docx", "url"]

SUPPORTED_EXTS = {
    ".txt": "txt",
    ".md": "md",
    ".markdown": "md",
    ".pdf": "pdf",
    ".docx": "docx",
}


class UnsupportedSourceError(ValueError):
    pass


def detect_kind(filename: str) -> Kind:
    name = filename.lower()
    for ext, kind in SUPPORTED_EXTS.items():
        if name.endswith(ext):
            return kind  # type: ignore[return-value]
    raise UnsupportedSourceError(f"unsupported file type: {filename}")


def extract_text(data: bytes, *, kind: Kind) -> str:
    if kind in ("txt", "md"):
        return data.decode("utf-8", errors="replace")
    if kind == "pdf":
        return _pdf_to_text(data)
    if kind == "docx":
        return _docx_to_text(data)
    raise UnsupportedSourceError(f"no extractor for kind={kind}")


def _pdf_to_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    out: list[str] = []
    for page in reader.pages:
        try:
            out.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001 — pypdf raises a grab-bag on malformed pages
            # Skip the page; downstream chunker will work with whatever did parse.
            from app.core.logging import log

            log.warning("kb.pdf.page_err", err=str(exc))
            continue
    return "\n\n".join(p for p in out if p.strip())


def _docx_to_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    paras = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    # Tables: flatten cell text row by row
    for tbl in doc.tables:
        for row in tbl.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text)
            if row_text:
                paras.append(row_text)
    return "\n\n".join(paras)
